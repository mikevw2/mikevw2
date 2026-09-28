"""Builds the 'All-in-One Budget & Debt Payoff Planner' workbook."""
import datetime as dt
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from common import (BOX, DARK, WHITE, DATE, GREY, INPUT, LIGHT, MID, MONEY, MONEY0, PCT, RED,
                    col, fill, header, section, style_range, title, widths)

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
INCOME_CATS = ["Salary", "Side Income", "Bonus", "Interest & Dividends", "Gifts", "Other Income"]
EXPENSE_CATS = [
    ("Rent / Mortgage", 1600), ("Utilities", 180), ("Internet & Phone", 120), ("Groceries", 550),
    ("Dining Out", 200), ("Transportation", 150), ("Car Payment", 350), ("Insurance", 220),
    ("Healthcare", 100), ("Debt Payments", 400), ("Subscriptions", 45), ("Shopping", 150),
    ("Personal Care", 60), ("Entertainment", 100), ("Travel", 150), ("Gifts & Donations", 75),
    ("Kids / Pets", 120), ("Savings", 500), ("Investing", 300), ("Miscellaneous", 80),
]
N_INC, N_EXP = 12, 30          # rows available for categories
TX_ROWS = 1500
N_DEBTS, SIM_MONTHS = 8, 360

SAMPLE_DEBTS = [("Visa Card", 4200, 0.2399, 120), ("Car Loan", 11800, 0.069, 350),
                ("Store Card", 850, 0.2699, 35), ("Student Loan", 18500, 0.0525, 210),
                ("Medical Bill", 1300, 0.0, 50)]
SAMPLE_GOALS = [("Emergency Fund", 10000, 3200, 18), ("Vacation", 3000, 650, 9),
                ("New Laptop", 1500, 400, 6), ("House Down Payment", 30000, 4200, 48)]


def build(path, sample=False, year=None, hook=None):
    year = year or dt.date.today().year
    wb = Workbook()
    start = wb.active
    start.title = "Start Here"
    setup = wb.create_sheet("Setup")
    tx = wb.create_sheet("Transactions")
    monthly = wb.create_sheet("Monthly Dashboard")
    annual = wb.create_sheet("Annual Overview")
    debt = wb.create_sheet("Debt Payoff")
    goals = wb.create_sheet("Savings Goals")
    calc = wb.create_sheet("DebtCalc")

    # ---------------- Start Here ----------------
    title(start, "Budget & Debt Payoff Planner", "Your whole financial year on one page — no formulas to write.", 12)
    steps = [
        ("1. Setup", "Enter the year, your income categories and a monthly budget for each expense category. Yellow cells are yours to edit."),
        ("2. Transactions", "Log every income and expense: date, description, type, category (drop-downs) and amount."),
        ("3. Monthly Dashboard", "Pick a month from the drop-down to see budget vs. actual, what's left, and your savings rate."),
        ("4. Annual Overview", "Every category across all 12 months — spot trends and seasonal spending instantly."),
        ("5. Debt Payoff", "List up to 8 debts, choose Snowball or Avalanche, add any extra payment and see your debt-free date."),
        ("6. Savings Goals", "Track up to 10 goals with the monthly amount needed to hit each target date."),
    ]
    r = 4
    for head, body in steps:
        start.cell(row=r, column=2, value=head).font = Font(bold=True, size=12, color=DARK)
        c = start.cell(row=r, column=3, value=body)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        start.row_dimensions[r].height = 36
        r += 1
    r += 1
    start.cell(row=r, column=2, value="Tips").font = Font(bold=True, size=12, color=DARK)
    tips = ["Only type in the yellow cells — everything else calculates automatically.",
            "Works in Microsoft Excel, Google Sheets (File → Import) and Apple Numbers.",
            "Rename any category on the Setup tab and every sheet updates automatically.",
            "Start a new year by saving a copy, clearing the Transactions tab and changing the year."]
    for t in tips:
        r += 1
        start.cell(row=r, column=3, value="• " + t)
    widths(start, {"B": 22, "C": 100})

    # ---------------- Setup ----------------
    title(setup, "Setup", "Yellow cells are inputs.", 9)
    setup["B4"] = "Budget year"
    setup["C4"] = year
    setup["B4"].font = Font(bold=True)
    setup["C4"].fill = fill(INPUT)
    setup["C4"].border = BOX
    section(setup, "B6", "Income categories")
    header(setup, 7, 2, ["Category"])
    for i in range(N_INC):
        c = setup.cell(row=8 + i, column=2, value=INCOME_CATS[i] if i < len(INCOME_CATS) else None)
        c.fill = fill(INPUT)
        c.border = BOX
    section(setup, "D6", "Expense categories")
    header(setup, 7, 4, ["Category", "Monthly budget"])
    for i in range(N_EXP):
        name, amt = EXPENSE_CATS[i] if i < len(EXPENSE_CATS) else (None, None)
        a = setup.cell(row=8 + i, column=4, value=name)
        b = setup.cell(row=8 + i, column=5, value=amt)
        for c in (a, b):
            c.fill = fill(INPUT)
            c.border = BOX
        b.number_format = MONEY
    setup.cell(row=8 + N_EXP, column=4, value="Total monthly budget").font = Font(bold=True)
    tot = setup.cell(row=8 + N_EXP, column=5, value=f"=SUM(E8:E{7 + N_EXP})")
    tot.number_format = MONEY
    tot.font = Font(bold=True)
    # helper lists (used by drop-downs)
    setup["H6"] = "Lists (used by drop-downs)"
    setup["H6"].font = Font(italic=True, color="888888")
    for i in range(N_INC):
        setup.cell(row=8 + i, column=8, value=f'=IF(B{8 + i}="","",B{8 + i})')
    for i in range(N_EXP):
        setup.cell(row=8 + N_INC + i, column=8, value=f'=IF(D{8 + i}="","",D{8 + i})')
    for i, m in enumerate(MONTHS):
        setup.cell(row=8 + i, column=9, value=m)
    setup.column_dimensions["H"].hidden = True
    setup.column_dimensions["I"].hidden = True
    widths(setup, {"B": 26, "C": 12, "D": 26, "E": 16, "F": 4})

    # ---------------- Transactions ----------------
    title(tx, "Transactions", "Log every income and expense. Month/Year fill in automatically.", 8)
    header(tx, 4, 1, ["Date", "Description", "Type", "Category", "Amount", "Month", "Year"])
    tx.freeze_panes = "A5"
    dv_type = DataValidation(type="list", formula1='"Income,Expense"', allow_blank=True)
    dv_cat = DataValidation(type="list", formula1=f"=Setup!$H$8:$H${7 + N_INC + N_EXP}", allow_blank=True)
    tx.add_data_validation(dv_type)
    tx.add_data_validation(dv_cat)
    last = 4 + TX_ROWS
    dv_type.add(f"C5:C{last}")
    dv_cat.add(f"D5:D{last}")
    for r in range(5, last + 1):
        tx.cell(row=r, column=1).number_format = DATE
        tx.cell(row=r, column=5).number_format = MONEY
        tx.cell(row=r, column=6, value=f'=IF(A{r}="","",MONTH(A{r}))').font = Font(color="999999")
        tx.cell(row=r, column=7, value=f'=IF(A{r}="","",YEAR(A{r}))').font = Font(color="999999")
        if r % 2 == 0:
            for c in range(1, 6):
                tx.cell(row=r, column=c).fill = fill(GREY)
    widths(tx, {"A": 14, "B": 34, "C": 11, "D": 24, "E": 14, "F": 8, "G": 8})
    if sample:
        _sample_transactions(tx, year)

    TXR = f"Transactions!$E$5:$E${last}"
    TYP = f"Transactions!$C$5:$C${last}"
    CAT = f"Transactions!$D$5:$D${last}"
    MON = f"Transactions!$F$5:$F${last}"
    YR = f"Transactions!$G$5:$G${last}"

    # ---------------- Monthly Dashboard ----------------
    title(monthly, "Monthly Dashboard", "Choose a month in the yellow cell.", 9)
    monthly["B4"] = "Month"
    monthly["B4"].font = Font(bold=True, size=12)
    monthly["C4"] = MONTHS[dt.date.today().month - 1] if not sample else "March"
    monthly["C4"].fill = fill(INPUT)
    monthly["C4"].font = Font(bold=True, size=12)
    monthly["C4"].border = BOX
    dv_m = DataValidation(type="list", formula1="=Setup!$I$8:$I$19", allow_blank=False)
    monthly.add_data_validation(dv_m)
    dv_m.add("C4")
    monthly["D4"] = "=Setup!$C$4"
    monthly["D4"].font = Font(bold=True, size=12)
    monthly["D4"].alignment = Alignment(horizontal="left")
    monthly["I4"] = "=MATCH(C4,Setup!$I$8:$I$19,0)"
    monthly["I4"].font = Font(color=WHITE)
    M = "$I$4"
    Y = "Setup!$C$4"

    # KPI tiles
    kpis = [("Income", f"=SUMIFS({TXR},{TYP},\"Income\",{MON},{M},{YR},{Y})", MONEY),
            ("Expenses", f"=SUMIFS({TXR},{TYP},\"Expense\",{MON},{M},{YR},{Y})", MONEY),
            ("Net cash flow", "=C7-D7", MONEY),
            ("Savings rate", '=IF(C7=0,0,E7/C7)', PCT),
            ("Budget left", f"=Setup!$E${8 + N_EXP}-D7", MONEY)]
    for i, (label, formula, fmt) in enumerate(kpis):
        cl = monthly.cell(row=6, column=3 + i, value=label)
        cl.font = Font(bold=True, color=WHITE)
        cl.fill = fill(MID)
        cl.alignment = Alignment(horizontal="center")
        cv = monthly.cell(row=7, column=3 + i, value=formula)
        cv.number_format = fmt
        cv.font = Font(bold=True, size=16, color=DARK)
        cv.fill = fill(LIGHT)
        cv.alignment = Alignment(horizontal="center")
    monthly.row_dimensions[7].height = 30

    section(monthly, "B9", "Expenses — budget vs. actual")
    header(monthly, 10, 2, ["Category", "Budget", "Actual", "Remaining", "% used"])
    er0 = 11
    for i in range(N_EXP):
        r = er0 + i
        monthly.cell(row=r, column=2, value=f'=IF(Setup!D{8 + i}="","",Setup!D{8 + i})')
        monthly.cell(row=r, column=3, value=f'=IF(B{r}="","",N(Setup!E{8 + i}))')
        monthly.cell(row=r, column=4, value=f'=IF(B{r}="","",SUMIFS({TXR},{TYP},"Expense",{CAT},B{r},{MON},{M},{YR},{Y}))')
        monthly.cell(row=r, column=5, value=f'=IF(B{r}="","",C{r}-D{r})')
        monthly.cell(row=r, column=6, value=f'=IF(OR(B{r}="",N(C{r})=0),"",D{r}/C{r})')
    erN = er0 + N_EXP - 1
    style_range(monthly, f"C{er0}:E{erN}", MONEY)
    style_range(monthly, f"B{er0}:B{erN}")
    style_range(monthly, f"F{er0}:F{erN}", "0%")
    monthly.conditional_formatting.add(f"E{er0}:E{erN}", CellIsRule(operator="lessThan", formula=["0"], fill=fill(RED)))
    monthly.conditional_formatting.add(f"F{er0}:F{erN}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E7D5B"))
    tr = erN + 1
    monthly.cell(row=tr, column=2, value="Total")
    monthly.cell(row=tr, column=3, value=f"=SUM(C{er0}:C{erN})")
    monthly.cell(row=tr, column=4, value=f"=SUM(D{er0}:D{erN})")
    monthly.cell(row=tr, column=5, value=f"=C{tr}-D{tr}")
    monthly.cell(row=tr, column=6, value=f'=IF(C{tr}=0,"",D{tr}/C{tr})')
    style_range(monthly, f"B{tr}:F{tr}", color=LIGHT, bold=True)
    style_range(monthly, f"C{tr}:E{tr}", MONEY, color=LIGHT, bold=True)
    monthly.cell(row=tr, column=6).number_format = "0%"

    section(monthly, "H9", "Income by category")
    header(monthly, 10, 8, ["Category", "Actual"])
    for i in range(N_INC):
        r = 11 + i
        monthly.cell(row=r, column=8, value=f'=IF(Setup!B{8 + i}="","",Setup!B{8 + i})')
        monthly.cell(row=r, column=9, value=f'=IF(H{r}="","",SUMIFS({TXR},{TYP},"Income",{CAT},H{r},{MON},{M},{YR},{Y}))')
    style_range(monthly, f"H11:H{10 + N_INC}")
    style_range(monthly, f"I11:I{10 + N_INC}", MONEY)

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Budget vs. actual"
    chart.style = 10
    chart.height = 11
    chart.width = 16
    n_named = len(EXPENSE_CATS)
    data = Reference(monthly, min_col=3, max_col=4, min_row=10, max_row=er0 + n_named - 1)
    cats = Reference(monthly, min_col=2, min_row=er0, max_row=er0 + n_named - 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.scaling.orientation = "minMax"
    chart.x_axis.scaling.orientation = "maxMin"
    monthly.add_chart(chart, f"H{12 + N_INC}")
    widths(monthly, {"B": 24, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 24, "I": 14})

    # ---------------- Annual Overview ----------------
    title(annual, "Annual Overview", "Every category, every month.", 16)
    annual["B3"] = "=\"Year: \"&Setup!$C$4"
    annual["B3"].font = Font(bold=True, color=WHITE)
    annual["B3"].fill = fill(DARK)
    heads = ["Category"] + [m[:3] for m in MONTHS] + ["Total", "Monthly avg"]
    header(annual, 5, 2, heads)
    annual.freeze_panes = "C6"

    def cat_block(first_row, setup_col, n, kind, label):
        for i in range(n):
            r = first_row + i
            annual.cell(row=r, column=2, value=f'=IF(Setup!{setup_col}{8 + i}="","",Setup!{setup_col}{8 + i})')
            for m in range(12):
                annual.cell(row=r, column=3 + m,
                            value=f'=IF($B{r}="","",SUMIFS({TXR},{TYP},"{kind}",{CAT},$B{r},{MON},{m + 1},{YR},{Y}))')
            annual.cell(row=r, column=15, value=f'=IF($B{r}="","",SUM(C{r}:N{r}))')
            annual.cell(row=r, column=16, value=f'=IF($B{r}="","",O{r}/12)')
        lastr = first_row + n - 1
        style_range(annual, f"B{first_row}:B{lastr}")
        style_range(annual, f"C{first_row}:P{lastr}", MONEY0)
        t = lastr + 1
        annual.cell(row=t, column=2, value=f"Total {label}")
        for c in range(3, 17):
            L = col(c)
            annual.cell(row=t, column=c, value=f"=SUM({L}{first_row}:{L}{lastr})")
        style_range(annual, f"B{t}:P{t}", color=LIGHT, bold=True)
        style_range(annual, f"C{t}:P{t}", MONEY0, color=LIGHT, bold=True)
        return t

    section(annual, "B6", "Income")
    inc_tot = cat_block(7, "B", N_INC, "Income", "income")
    section(annual, f"B{inc_tot + 2}", "Expenses")
    exp_tot = cat_block(inc_tot + 3, "D", N_EXP, "Expense", "expenses")
    net = exp_tot + 2
    annual.cell(row=net, column=2, value="Net cash flow")
    annual.cell(row=net + 1, column=2, value="Savings rate")
    for c in range(3, 17):
        L = col(c)
        annual.cell(row=net, column=c, value=f"={L}{inc_tot}-{L}{exp_tot}")
        annual.cell(row=net + 1, column=c, value=f'=IF({L}{inc_tot}=0,"",{L}{net}/{L}{inc_tot})')
    style_range(annual, f"B{net}:P{net}", color=DARK, bold=True)
    style_range(annual, f"C{net}:P{net}", MONEY0, color=DARK, bold=True)
    for c in range(2, 17):
        annual.cell(row=net, column=c).font = Font(bold=True, color=WHITE)
    style_range(annual, f"B{net + 1}:B{net + 1}", bold=True)
    style_range(annual, f"C{net + 1}:P{net + 1}", PCT)
    lc = LineChart()
    lc.title = "Income vs. expenses by month"
    lc.height = 8
    lc.width = 24
    for rr in (inc_tot, exp_tot):
        lc.add_data(Reference(annual, min_col=2, max_col=14, min_row=rr), from_rows=True, titles_from_data=True)
    lc.set_categories(Reference(annual, min_col=3, max_col=14, min_row=5))
    annual.add_chart(lc, f"B{net + 3}")
    widths(annual, {"B": 24, **{col(c): 10 for c in range(3, 15)}, "O": 12, "P": 12})

    # ---------------- Debt Payoff ----------------
    title(debt, "Debt Payoff Planner", "Snowball (smallest balance first) or Avalanche (highest interest first).", 10)
    header(debt, 4, 2, ["Debt name", "Balance", "APR", "Min. payment", "Paid off in (months)", "Payoff date"])
    for i in range(N_DEBTS):
        r = 5 + i
        vals = SAMPLE_DEBTS[i] if sample and i < len(SAMPLE_DEBTS) else (None,) * 4
        for j, v in enumerate(vals):
            c = debt.cell(row=r, column=2 + j, value=v)
            c.fill = fill(INPUT)
            c.border = BOX
        debt.cell(row=r, column=3).number_format = MONEY
        debt.cell(row=r, column=4).number_format = "0.00%"
        debt.cell(row=r, column=5).number_format = MONEY
        debt.cell(row=r, column=6, value=f'=IF(N(C{r})=0,"",INDEX(DebtCalc!$F$12:$F$19,DebtCalc!$C${2 + i}))')
        debt.cell(row=r, column=7, value=f'=IF(OR(F{r}="",F{r}>{SIM_MONTHS}),"",EDATE($J$7,F{r}))')
        debt.cell(row=r, column=7).number_format = "mmm yyyy"
        style_range(debt, f"F{r}:G{r}")
    tr = 5 + N_DEBTS
    debt.cell(row=tr, column=2, value="Total")
    debt.cell(row=tr, column=3, value=f"=SUM(C5:C{tr - 1})").number_format = MONEY
    debt.cell(row=tr, column=5, value=f"=SUM(E5:E{tr - 1})").number_format = MONEY
    style_range(debt, f"B{tr}:G{tr}", color=LIGHT, bold=True)
    debt.cell(row=tr, column=3).number_format = MONEY
    debt.cell(row=tr, column=5).number_format = MONEY

    labels = [("Extra payment each month", 200 if sample else 0, MONEY),
              ("Strategy", "Snowball", None),
              ("Start date", dt.date(year, 1, 1) if sample else dt.date.today().replace(day=1), DATE)]
    for i, (lab, val, fmt) in enumerate(labels):
        debt.cell(row=5 + i, column=9, value=lab).font = Font(bold=True)
        c = debt.cell(row=5 + i, column=10, value=val)
        c.fill = fill(INPUT)
        c.border = BOX
        if fmt:
            c.number_format = fmt
    dv_s = DataValidation(type="list", formula1='"Snowball,Avalanche"', allow_blank=False)
    debt.add_data_validation(dv_s)
    dv_s.add("J6")

    section(debt, "I9", "Your results")
    SIM_FIRST, SIM_LAST = 25, 24 + SIM_MONTHS
    results = [
        ("Total paid toward debt each month", f"=E{tr}+J5", MONEY),
        ("Months until debt-free", f'=IF(C{tr}=0,0,IF(DebtCalc!$AA${SIM_LAST}>0.005,"30+ years — raise payments",COUNTIF(DebtCalc!$AA${SIM_FIRST}:$AA${SIM_LAST},">0.005")+1))', "0"),
        ("Debt-free date", f'=IF(ISNUMBER(J11),EDATE(J7,J11),"—")', "mmmm yyyy"),
        ("Total interest paid", f"=IF(ISNUMBER(J11),DebtCalc!$AC$2,\"—\")", MONEY),
        ("Total of all payments", f"=IF(ISNUMBER(J11),DebtCalc!$AC$3,\"—\")", MONEY),
    ]
    for i, (lab, f, fmt) in enumerate(results):
        debt.cell(row=10 + i, column=9, value=lab)
        c = debt.cell(row=10 + i, column=10, value=f)
        c.number_format = fmt
        c.font = Font(bold=True, size=12, color=DARK)
        c.fill = fill(LIGHT)
        c.border = BOX
    debt["I16"] = "Tip: try adding $50/month extra and watch the date move."
    debt["I16"].font = Font(italic=True, color="666666")
    lc2 = LineChart()
    lc2.title = "Total debt remaining"
    lc2.height = 8
    lc2.width = 20
    lc2.add_data(Reference(calc, min_col=27, min_row=24, max_row=24 + 120), titles_from_data=False)
    lc2.legend = None
    lc2.y_axis.title = "Balance"
    lc2.x_axis.title = "Months from start (first 10 years)"
    debt.add_chart(lc2, "B16")
    widths(debt, {"B": 20, "C": 14, "D": 10, "E": 14, "F": 12, "G": 12, "H": 3, "I": 34, "J": 24})

    # ---------------- DebtCalc (hidden engine) ----------------
    calc["A1"] = "Engine for the Debt Payoff tab. Do not edit."
    for i in range(N_DEBTS):
        r = 2 + i
        dr = 5 + i
        calc.cell(row=r, column=1, value=i + 1)
        calc.cell(row=r, column=2, value=(f"=IF(N('Debt Payoff'!C{dr})=0,1E+12+A{r},"
                                          f"IF('Debt Payoff'!$J$6=\"Avalanche\",-N('Debt Payoff'!D{dr}),N('Debt Payoff'!C{dr}))+A{r}/1E+6)"))
        calc.cell(row=r, column=3, value=f'=COUNTIF($B$2:$B$9,"<"&B{r})+1')
    for k in range(N_DEBTS):
        r = 12 + k
        calc.cell(row=r, column=1, value=k + 1)
        calc.cell(row=r, column=2, value=f"=MATCH(SMALL($B$2:$B$9,A{r}),$B$2:$B$9,0)")
        calc.cell(row=r, column=3, value=f"=N(INDEX('Debt Payoff'!$C$5:$C$12,B{r}))")
        calc.cell(row=r, column=4, value=f"=N(INDEX('Debt Payoff'!$D$5:$D$12,B{r}))/12")
        calc.cell(row=r, column=5, value=f"=N(INDEX('Debt Payoff'!$E$5:$E$12,B{r}))")
        endc = col(4 + 3 * k)
        calc.cell(row=r, column=6, value=f'=IF(C{r}=0,"",COUNTIF({endc}${SIM_FIRST}:{endc}${SIM_LAST},">0.005")+1)')
    # simulation grid: per debt k -> MinPay, Extra, End ; Z leftover, AA total balance
    hdr = ["Month"]
    for k in range(N_DEBTS):
        hdr += [f"D{k + 1} min", f"D{k + 1} extra", f"D{k + 1} end"]
    hdr += ["Leftover", "Total balance"]
    for j, h in enumerate(hdr):
        calc.cell(row=23, column=1 + j, value=h)
    calc.cell(row=24, column=1, value=0)
    for k in range(N_DEBTS):
        calc.cell(row=24, column=4 + 3 * k, value=f"=C{12 + k}")
    calc.cell(row=24, column=27, value="=" + "+".join(f"{col(4 + 3 * k)}24" for k in range(N_DEBTS)))
    for m in range(1, SIM_MONTHS + 1):
        r = 24 + m
        calc.cell(row=r, column=1, value=m)
        for k in range(N_DEBTS):
            mc, ec, bc = col(2 + 3 * k), col(3 + 3 * k), col(4 + 3 * k)
            owed = f"ROUND({bc}{r - 1}*(1+$D${12 + k}),2)"
            calc[f"{mc}{r}"] = f"=MIN($E${12 + k},{owed})"
            prior = "+".join(f"{col(3 + 3 * j)}{r}" for j in range(k)) or "0"
            calc[f"{ec}{r}"] = f"=MAX(0,MIN({owed}-{mc}{r},$Z{r}-({prior})))"
            calc[f"{bc}{r}"] = f"=ROUND({owed}-{mc}{r}-{ec}{r},2)"
        mins = "+".join(f"{col(2 + 3 * k)}{r}" for k in range(N_DEBTS))
        calc[f"Z{r}"] = f"=MAX(0,'Debt Payoff'!$J$10-({mins}))"
        calc[f"AA{r}"] = "=" + "+".join(f"{col(4 + 3 * k)}{r}" for k in range(N_DEBTS))
    pay_cols = ",".join(f"{col(2 + 3 * k)}{SIM_FIRST}:{col(3 + 3 * k)}{SIM_LAST}" for k in range(N_DEBTS))
    calc["AB3"] = "Total paid"
    calc["AC3"] = f"=SUM({pay_cols})"
    calc["AB2"] = "Total interest"
    calc["AC2"] = "=AC3-SUM(C12:C19)"
    calc.sheet_state = "hidden"

    # ---------------- Savings Goals ----------------
    title(goals, "Savings Goals", "Set a target and a date — see exactly what to save each month.", 10)
    header(goals, 4, 2, ["Goal", "Target", "Saved so far", "Target date", "Remaining", "Months left", "Save per month", "Progress"])
    today = dt.date.today()
    for i in range(10):
        r = 5 + i
        if sample and i < len(SAMPLE_GOALS):
            g, t, s, months = SAMPLE_GOALS[i]
            yy, mm = divmod(today.month - 1 + months, 12)
            vals = [g, t, s, dt.date(today.year + yy, mm + 1, 1)]
        else:
            vals = [None] * 4
        for j, v in enumerate(vals):
            c = goals.cell(row=r, column=2 + j, value=v)
            c.fill = fill(INPUT)
            c.border = BOX
        goals.cell(row=r, column=3).number_format = MONEY
        goals.cell(row=r, column=4).number_format = MONEY
        goals.cell(row=r, column=5).number_format = DATE
        goals.cell(row=r, column=6, value=f'=IF(B{r}="","",MAX(0,C{r}-D{r}))').number_format = MONEY
        goals.cell(row=r, column=7, value=f'=IF(OR(B{r}="",E{r}=""),"",MAX(1,DATEDIF(TODAY(),E{r},"m")))')
        goals.cell(row=r, column=8, value=f'=IF(OR(F{r}="",G{r}=""),"",F{r}/G{r})').number_format = MONEY
        goals.cell(row=r, column=9, value=f'=IF(OR(B{r}="",N(C{r})=0),"",MIN(1,D{r}/C{r}))').number_format = "0%"
        style_range(goals, f"F{r}:I{r}")
    goals.conditional_formatting.add("I5:I14", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E7D5B"))
    widths(goals, {"B": 24, "C": 14, "D": 14, "E": 14, "F": 14, "G": 12, "H": 16, "I": 22})

    for ws in (setup, tx, debt, goals):
        ws.sheet_properties.tabColor = "F2C94C"
    for ws in (monthly, annual):
        ws.sheet_properties.tabColor = MID
    wb.active = 0
    if hook:
        hook(wb)
    wb.save(path)





def _sample_transactions(tx, year):
    rnd = random.Random(42)
    rows = []
    for m in range(1, 13):
        rows.append((dt.date(year, m, 1), "Paycheck", "Income", "Salary", 2650.00))
        rows.append((dt.date(year, m, 15), "Paycheck", "Income", "Salary", 2650.00))
        if m % 2 == 0:
            rows.append((dt.date(year, m, 20), "Etsy shop payout", "Income", "Side Income", round(rnd.uniform(180, 520), 2)))
        fixed = [("Rent", "Rent / Mortgage", 1600), ("Electric & water", "Utilities", rnd.uniform(140, 210)),
                 ("Phone + internet", "Internet & Phone", 118.5), ("Car payment", "Car Payment", 350),
                 ("Auto + renters insurance", "Insurance", 214), ("Credit card payments", "Debt Payments", 405),
                 ("Netflix / Spotify / iCloud", "Subscriptions", 42.97), ("Transfer to savings", "Savings", 500),
                 ("Roth IRA", "Investing", 300)]
        for d, c, a in fixed:
            rows.append((dt.date(year, m, rnd.randint(1, 5)), d, "Expense", c, round(a, 2)))
        for _ in range(5):
            rows.append((dt.date(year, m, rnd.randint(1, 28)), rnd.choice(["Kroger", "Trader Joe's", "Costco", "Aldi"]),
                         "Expense", "Groceries", round(rnd.uniform(60, 140), 2)))
        for _ in range(rnd.randint(3, 6)):
            rows.append((dt.date(year, m, rnd.randint(1, 28)), rnd.choice(["Chipotle", "Pizza night", "Coffee shop", "Brunch"]),
                         "Expense", "Dining Out", round(rnd.uniform(12, 70), 2)))
        for _ in range(3):
            rows.append((dt.date(year, m, rnd.randint(1, 28)), "Gas station", "Expense", "Transportation", round(rnd.uniform(35, 58), 2)))
        rows.append((dt.date(year, m, rnd.randint(1, 28)), "Target run", "Expense", "Shopping", round(rnd.uniform(40, 220), 2)))
        rows.append((dt.date(year, m, rnd.randint(1, 28)), "Movies / concert", "Expense", "Entertainment", round(rnd.uniform(20, 140), 2)))
        rows.append((dt.date(year, m, rnd.randint(1, 28)), "Haircut", "Expense", "Personal Care", 45.0))
        rows.append((dt.date(year, m, rnd.randint(1, 28)), "Vet / pet food", "Expense", "Kids / Pets", round(rnd.uniform(40, 160), 2)))
        if m in (6, 7, 12):
            rows.append((dt.date(year, m, 10), "Trip", "Expense", "Travel", round(rnd.uniform(350, 900), 2)))
    rows.sort(key=lambda x: x[0])
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tx.cell(row=5 + i, column=1 + j, value=v)


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    build(f"{out}/Budget-and-Debt-Payoff-Planner.xlsx", sample=False)
    build(f"{out}/Budget-and-Debt-Payoff-Planner-SAMPLE.xlsx", sample=True, year=2026)
