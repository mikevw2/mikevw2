"""Shared styling helpers for the spreadsheet products."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DARK = "1F4E3D"
MID = "2E7D5B"
LIGHT = "E8F3EC"
INPUT = "FFF8E1"
GREY = "F4F5F6"
WHITE = "FFFFFF"
RED = "F8D7DA"

MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
PCT = "0.0%"
DATE = "mmm d, yyyy"

thin = Side(style="thin", color="D0D7D3")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def title(ws, text, subtitle=None, width_cols=10):
    ws.sheet_view.showGridLines = False
    for c in range(1, width_cols + 1):
        ws.cell(row=1, column=c).fill = fill(DARK)
        ws.cell(row=2, column=c).fill = fill(DARK)
    ws["B1"] = text
    ws["B1"].font = Font(name="Calibri", size=20, bold=True, color=WHITE)
    ws.row_dimensions[1].height = 34
    if subtitle:
        ws["B2"] = subtitle
        ws["B2"].font = Font(name="Calibri", size=11, italic=True, color="CFE3D8")
    ws.column_dimensions["A"].width = 3


def header(ws, row, col, labels, color=MID):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font = Font(bold=True, color=WHITE)
        c.fill = fill(color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def section(ws, cell, text):
    ws[cell] = text
    ws[cell].font = Font(size=13, bold=True, color=DARK)


def style_range(ws, rng, fmt=None, color=None, bold=False, border=True):
    for row in ws[rng]:
        for c in row:
            if fmt:
                c.number_format = fmt
            if color:
                c.fill = fill(color)
            if bold:
                c.font = Font(bold=True)
            if border:
                c.border = BOX


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def col(n):
    return get_column_letter(n)
