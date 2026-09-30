"""PRO-tier tabs for the Food Safety Audit-Readiness Kit.

add_pro_tabs() appends six tools to a workbook built by food_safety_kit.build(pro=True)
and returns the dashboard tiles and action-list rows they contribute.
"""
import datetime as dt
import random

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font

from common import BOX, DATE, INPUT, PCT, RED, col, fill, header, section
from food_safety_kit import AMBER, _dv, _log, _put, _status_colors

ALLERGENS = ["Milk", "Egg", "Fish", "Crustacean shellfish", "Tree nuts", "Peanuts", "Wheat", "Soy", "Sesame"]
DEVICES = [("Rodent trap", 2), ("Bait station", 1), ("Insect light trap", 20), ("Pheromone trap", 5)]
INSTRUMENTS = ["Thermometer", "Recording thermometer", "pH meter", "Scale", "Metal detector",
               "Pressure gauge", "Test kit / titration", "Other"]

TAB_NOTES = [
    ("Mock Recall", "Traceability exercises with time to complete and % of product accounted for, scored PASS / FAIL against your targets."),
    ("Allergen Matrix", "Products × the 9 major US allergens, with automatic 'Contains' and 'Shared line' statements, label verification status and a line allergen profile."),
    ("Glass Register", "Glass and brittle-plastic items by location, with inspection schedule, condition and DAMAGED / OVERDUE flags."),
    ("Pest Trend", "Pest device log with a 12-month trend by device type. Months over your action threshold turn red."),
    ("Calibration Log", "Thermometers, scales, pH meters and metal detectors with calibration due dates and failed calibrations flagged."),
    ("Internal Audits", "Annual internal audit schedule covering every program area, with coverage % for the last 12 months."),
]


def add_pro_tabs(wb, ctx):
    today, sample = ctx["today"], ctx["sample"]
    WINDOW, FREQ_TABLE, FREQ_LIST, DEPT_LIST = ctx["WINDOW"], ctx["FREQ_TABLE"], ctx["FREQ_LIST"], ctx["DEPT_LIST"]
    due_soon = f"MIN({WINDOW},VLOOKUP({{f}},{FREQ_TABLE},2,FALSE)/4)"
    kpis, actions = [], []

    # ---------------- Mock Recall ----------------
    mr = wb.create_sheet("Mock Recall")
    mr_cols = [("Date", 12, None, DATE), ("Type", 11, None, None), ("Product / lot traced", 24, None, None),
               ("Scenario", 26, None, None), ("Start", 15, None, "mm/dd hh:mm"), ("End", 15, None, "mm/dd hh:mm"),
               ("Hours", 8, '=IF(OR(E{r}="",F{r}=""),"",(F{r}-E{r})*24)', "0.0"),
               ("Qty produced", 11, None, "#,##0"), ("Qty located", 11, None, "#,##0"),
               ("% accounted", 10, '=IF(OR(H{r}="",I{r}="",N(H{r})=0),"",I{r}/H{r})', PCT),
               ("Result", 12, '=IF(A{r}="","",IF(COUNT(E{r},F{r},H{r},I{r})<4,"INCOMPLETE",IF(AND(G{r}<=$D$4,J{r}>=$D$5),"PASS","FAIL")))', None),
               ("Gaps found", 28, None, None), ("CAPA #", 9, None, None)]
    mr_last = _log(mr, "Mock Recall & Traceability", "Run at least once per exercise period. Trace one step forward and one step back.",
                   mr_cols, first=10, n=60)
    for i, (lab, v, fmt) in enumerate([("Target: time to complete (hours)", 4, "0.0"),
                                       ("Target: product accounted for", 1.0, "0%"),
                                       ("Exercise required every (months)", 12, "0")]):
        mr.cell(row=4 + i, column=1, value=lab).font = Font(bold=True)
        mr.merge_cells(start_row=4 + i, start_column=1, end_row=4 + i, end_column=3)
        c = mr.cell(row=4 + i, column=4, value=v)
        c.fill = fill(INPUT)
        c.border = BOX
        c.number_format = fmt
    mr.freeze_panes = "B10"
    _dv(mr, '"Forward,Backward,Both"', f"B10:B{mr_last}")
    _status_colors(mr, f"K10:K{mr_last}", bad=["FAIL"], warn=["INCOMPLETE"], good=["PASS"])
    dates = f"'Mock Recall'!$A$10:$A${mr_last}"
    kpis += [("Days since mock recall", f'=IF(COUNT({dates})=0,"None",TODAY()-MAX({dates}))', "0", ("gt", "'Mock Recall'!$D$6*30.44")),
             ("Last mock recall", f'=IF(COUNT({dates})=0,"—",INDEX(\'Mock Recall\'!$K$10:$K${mr_last},MATCH(MAX({dates}),{dates},0)))', None, ("eq", '"FAIL"'))]
    actions.append(("Mock recall", "Mock Recall", "FAIL / overdue"))

    # ---------------- Internal Audits ----------------
    ia = wb.create_sheet("Internal Audits")
    ia_cols = [("Program area", 34, None, None), ("Lead auditor", 16, None, None), ("Planned date", 13, None, DATE),
               ("Completed", 13, None, DATE), ("Findings", 9, None, "0"), ("CAPA #s", 18, None, None),
               ("Status", 14, '=IF(A{r}="","",IF(D{r}<>"",IF(TODAY()-D{r}>365,"DUE AGAIN","COMPLETE"),IF(C{r}="","NOT SCHEDULED",'
                              'IF(C{r}<TODAY(),"OVERDUE",IF(C{r}-TODAY()<=30,"DUE SOON","PLANNED")))))', None)]
    ia_last = _log(ia, "Internal Audit Schedule", "Every program area audited at least once every 12 months, by someone independent of the area.",
                   ia_cols, n=40)
    for i, sec in enumerate(ctx["sections"]):
        ia.cell(row=5 + i, column=1, value=sec)
    _status_colors(ia, f"G5:G{ia_last}", bad=["OVERDUE", "DUE AGAIN"], warn=["DUE SOON", "NOT SCHEDULED"], good=["COMPLETE"])
    ia_area = f"'Internal Audits'!$A$5:$A${ia_last}"
    ia_done = f"'Internal Audits'!$D$5:$D${ia_last}"
    ia_status = f"'Internal Audits'!$G$5:$G${ia_last}"
    kpis += [("Internal audit coverage", f'=IFERROR(COUNTIFS({ia_done},">="&(TODAY()-365))/COUNTA({ia_area}),"—")', PCT, None),
             ("Internal audits overdue", f'=COUNTIF({ia_status},"OVERDUE")+COUNTIF({ia_status},"DUE AGAIN")', "0", ("gt", "0"))]
    actions.append(("Internal audits", "Internal Audits", "OVERDUE"))

    # ---------------- Allergen Matrix ----------------
    al = wb.create_sheet("Allergen Matrix")
    n_al = len(ALLERGENS)
    c0 = 3  # first allergen column (C)
    last_al = c0 + n_al - 1
    L = col

    def statement(r, word):
        parts = "&".join(f'IF({L(c0 + j)}{r}="{word}",{L(c0 + j)}$4&", ","")' for j in range(n_al))
        return f'IF($A{r}="","",IF(LEN({parts})=0,"None",LEFT({parts},LEN({parts})-2)))'

    al_cols = ([("Product", 24, None, None), ("Line", 12, None, None)]
               + [(a, 12 if len(a) > 10 else 9, None, None) for a in ALLERGENS]
               + [("Contains statement", 30, "=" + statement("{r}", "Contains"), None),
                  ("Shared-line / may contain", 24, "=" + statement("{r}", "Shared line"), None),
                  ("Label verified", 12, None, DATE),
                  ("Label status", 13, f'=IF(A{{r}}="","",IF({L(last_al + 3)}{{r}}="","NOT VERIFIED",IF(TODAY()-{L(last_al + 3)}{{r}}>365,"RE-VERIFY","OK")))', None)])
    al_last = _log(al, "Allergen Matrix", "Mark each allergen 'Contains' or 'Shared line'. Statements build themselves. Re-verify labels at least yearly.",
                   al_cols, n=100)
    al.row_dimensions[4].height = 44
    _dv(al, '"Contains,Shared line"', f"{L(c0)}5:{L(last_al)}{al_last}")
    al.conditional_formatting.add(f"{L(c0)}5:{L(last_al)}{al_last}", CellIsRule(operator="equal", formula=['"Contains"'], fill=fill(RED)))
    al.conditional_formatting.add(f"{L(c0)}5:{L(last_al)}{al_last}", CellIsRule(operator="equal", formula=['"Shared line"'], fill=fill(AMBER)))
    status_col = L(last_al + 4)
    _status_colors(al, f"{status_col}5:{status_col}{al_last}", bad=["NOT VERIFIED", "RE-VERIFY"], warn=[], good=["OK"])
    # line allergen profile
    pc = last_al + 6
    section(al, f"{L(pc)}3", "Line allergen profile (● = an allergen runs on this line)")
    header(al, 4, pc, ["Line"] + ALLERGENS)
    for i in range(10):
        r = 5 + i
        c = al.cell(row=r, column=pc)
        c.fill = fill(INPUT)
        c.border = BOX
        for j in range(n_al):
            ac = L(c0 + j)
            cell = al.cell(row=r, column=pc + 1 + j,
                           value=f'=IF({L(pc)}{r}="","",IF(COUNTIFS($B$5:$B${al_last},{L(pc)}{r},{ac}$5:{ac}${al_last},"Contains")>0,"●",""))')
            cell.alignment = Alignment(horizontal="center")
            cell.border = BOX
    al.column_dimensions[L(pc)].width = 14
    for j in range(n_al):
        al.column_dimensions[L(pc + 1 + j)].width = 9
    al.column_dimensions[L(pc - 1)].width = 3
    al_status = f"'Allergen Matrix'!${status_col}$5:${status_col}${al_last}"
    kpis.append(("Allergen labels to verify", f'=COUNTIF({al_status},"NOT VERIFIED")+COUNTIF({al_status},"RE-VERIFY")', "0", ("gt", "0")))
    actions.append(("Allergen labels", "Allergen Matrix", "RE-VERIFY"))

    # ---------------- Glass Register ----------------
    gl = wb.create_sheet("Glass Register")
    gl_cols = [("Item ID", 9, None, None), ("Department", 16, None, None), ("Location", 20, None, None),
               ("Description", 24, None, None), ("Material", 14, None, None), ("Qty", 6, None, "0"),
               ("Frequency", 11, None, None), ("Last inspected", 13, None, DATE), ("Condition", 10, None, None),
               ("Inspected by", 12, None, None),
               ("Next due", 12, f'=IF(OR(H{{r}}="",G{{r}}=""),"",H{{r}}+VLOOKUP(G{{r}},{FREQ_TABLE},2,FALSE))', DATE),
               ("Status", 14, '=IF(A{r}="","",IF(OR(I{r}="Damaged",I{r}="Missing"),"DAMAGED",IF(K{r}="","NOT INSPECTED",'
                              'IF(K{r}<TODAY(),"OVERDUE",IF(K{r}-TODAY()<=' + due_soon.format(f="G{r}") + ',"DUE SOON","OK")))))', None)]
    gl_last = _log(gl, "Glass & Brittle Plastic Register", "Any break is an incident: stop, isolate, clean up, inspect product, record it and open a CAPA.", gl_cols)
    _dv(gl, DEPT_LIST, f"B5:B{gl_last}")
    _dv(gl, '"Glass,Brittle plastic,Acrylic,Ceramic,Other"', f"E5:E{gl_last}")
    _dv(gl, FREQ_LIST, f"G5:G{gl_last}")
    _dv(gl, '"Intact,Damaged,Missing"', f"I5:I{gl_last}")
    _status_colors(gl, f"L5:L{gl_last}", bad=["DAMAGED", "OVERDUE"], warn=["DUE SOON", "NOT INSPECTED"], good=["OK"])
    gl_status = f"'Glass Register'!$L$5:$L${gl_last}"
    kpis += [("Glass / plastic issues", f'=COUNTIF({gl_status},"DAMAGED")+COUNTIF({gl_status},"OVERDUE")', "0", ("gt", "0")),
             ("Glass checks due soon", f'=COUNTIF({gl_status},"DUE SOON")', "0", None)]
    actions.append(("Glass & brittle plastic", "Glass Register", "DAMAGED / OVERDUE"))

    # ---------------- Pest Trend ----------------
    pt = wb.create_sheet("Pest Trend")
    pt_cols = [("Date", 12, None, DATE), ("Device ID", 9, None, None), ("Device type", 16, None, None),
               ("Area", 18, None, None), ("Pest", 14, None, None), ("Count", 7, None, "0"),
               ("Action taken", 26, None, None), ("Technician", 12, None, None)]
    pt_last = _log(pt, "Pest Activity & Trend", "Log every device with activity (count 0 rows are optional). The trend table flags months over threshold.", pt_cols, n=500)
    tc = 10  # trend table starts at column J
    section(pt, f"{L(tc)}3", "12-month trend (activity count)")
    header(pt, 4, tc, ["Month"] + [d for d, _ in DEVICES] + ["Alert"])
    pt.row_dimensions[4].height = 32
    pt.cell(row=5, column=tc, value="Action threshold").font = Font(bold=True)
    pt.cell(row=5, column=tc).border = BOX
    for j, (_, thr) in enumerate(DEVICES):
        c = pt.cell(row=5, column=tc + 1 + j, value=thr)
        c.fill = fill(INPUT)
        c.border = BOX
    first_dev, last_dev = L(tc + 1), L(tc + len(DEVICES))
    alert_col = L(tc + len(DEVICES) + 1)
    for i in range(12):
        r = 6 + i
        m = pt.cell(row=r, column=tc, value=f"=EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),{i - 11})")
        m.number_format = "mmm yyyy"
        m.border = BOX
        for j in range(len(DEVICES)):
            dc = L(tc + 1 + j)
            c = pt.cell(row=r, column=tc + 1 + j,
                        value=f'=SUMIFS($F$5:$F${pt_last},$C$5:$C${pt_last},{dc}$4,$A$5:$A${pt_last},">="&${L(tc)}{r},$A$5:$A${pt_last},"<"&EDATE(${L(tc)}{r},1))')
            c.border = BOX
            c.alignment = Alignment(horizontal="center")
        a = pt.cell(row=r, column=tc + len(DEVICES) + 1,
                    value=f'=IF(SUMPRODUCT(--({first_dev}{r}:{last_dev}{r}>${first_dev}$5:${last_dev}$5))>0,"ALERT","")')
        a.border = BOX
        a.font = Font(bold=True, color="9B1C1C")
    pt.conditional_formatting.add(f"{first_dev}6:{last_dev}17", FormulaRule(formula=[f"{first_dev}6>{first_dev}$5"], fill=fill(RED)))
    _dv(pt, f"=${first_dev}$4:${last_dev}$4", f"C5:C{pt_last}")
    _dv(pt, '"Interior - processing,Interior - warehouse,Interior - other,Exterior"', f"D5:D{pt_last}")
    _dv(pt, '"Mouse,Rat,Flies,Stored-product insects,Cockroach,Birds,Other"', f"E5:E{pt_last}")
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.title = "Pest activity by month"
    ch.height = 8
    ch.width = 17
    for j, (name, _) in enumerate(DEVICES):  # months only (rows 6-17), skipping the threshold row
        ch.add_data(Reference(pt, min_col=tc + 1 + j, min_row=6, max_row=17), titles_from_data=False)
        ch.series[-1].tx = SeriesLabel(v=name)
    ch.set_categories(Reference(pt, min_col=tc, min_row=6, max_row=17))
    pt.add_chart(ch, f"{L(tc)}19")
    for j in range(len(DEVICES) + 2):
        pt.column_dimensions[L(tc + j)].width = 12
    pt.column_dimensions[L(tc)].width = 16
    pt.column_dimensions["I"].width = 3
    kpis += [("Pest alerts (this month)", f"=SUMPRODUCT(--('Pest Trend'!${first_dev}$17:${last_dev}$17>'Pest Trend'!${first_dev}$5:${last_dev}$5))", "0", ("gt", "0")),
             ("Pest activity (this month)", f"=SUM('Pest Trend'!${first_dev}$17:${last_dev}$17)", "0", None)]
    actions.append(("Pest trend", "Pest Trend", "red months"))

    # ---------------- Calibration Log ----------------
    ca = wb.create_sheet("Calibration Log")
    ca_cols = [("Instrument ID", 11, None, None), ("Description", 24, None, None), ("Location", 16, None, None),
               ("Type", 18, None, None), ("Reference / method", 22, None, None), ("Tolerance", 11, None, None),
               ("Frequency", 11, None, None), ("Last calibrated", 13, None, DATE), ("Result", 10, None, None),
               ("Performed by", 13, None, None),
               ("Next due", 12, f'=IF(OR(H{{r}}="",G{{r}}=""),"",H{{r}}+VLOOKUP(G{{r}},{FREQ_TABLE},2,FALSE))', DATE),
               ("Status", 15, '=IF(A{r}="","",IF(I{r}="Fail","FAILED",IF(K{r}="","NOT CALIBRATED",'
                              'IF(K{r}<TODAY(),"OVERDUE",IF(K{r}-TODAY()<=' + due_soon.format(f="G{r}") + ',"DUE SOON","OK")))))', None)]
    ca_last = _log(ca, "Calibration Log", "A failed calibration means product since the last good check may be affected. Evaluate it and open a CAPA.", ca_cols, n=200)
    _dv(ca, '"' + ",".join(INSTRUMENTS) + '"', f"D5:D{ca_last}")
    _dv(ca, FREQ_LIST, f"G5:G{ca_last}")
    _dv(ca, '"Pass,Adjusted,Fail"', f"I5:I{ca_last}")
    _status_colors(ca, f"L5:L{ca_last}", bad=["FAILED", "OVERDUE"], warn=["DUE SOON", "NOT CALIBRATED"], good=["OK"])
    ca_status = f"'Calibration Log'!$L$5:$L${ca_last}"
    kpis.append(("Calibration overdue / failed", f'=COUNTIF({ca_status},"OVERDUE")+COUNTIF({ca_status},"FAILED")', "0", ("gt", "0")))
    actions.append(("Calibration", "Calibration Log", "OVERDUE / FAILED"))

    for ws in (mr, ia, al, gl, pt, ca):
        ws.sheet_properties.tabColor = "7B61A8"
    if sample:
        _sample(today, mr, ia, al, gl, pt, ca, ctx["sections"])
    # tile order on the dashboard (two rows of five)
    order = ["Days since mock recall", "Last mock recall", "Internal audit coverage", "Internal audits overdue",
             "Allergen labels to verify", "Glass / plastic issues", "Glass checks due soon", "Pest alerts (this month)",
             "Pest activity (this month)", "Calibration overdue / failed"]
    kpis.sort(key=lambda k: order.index(k[0]))
    return kpis, actions


def _sample(today, mr, ia, al, gl, pt, ca, sections):
    rnd = random.Random(23)
    d = lambda n: today - dt.timedelta(days=n)  # noqa: E731
    at = lambda n, h, m: dt.datetime.combine(d(n), dt.time(h, m))  # noqa: E731
    _put(mr, 10, [d(390), "Both", "Whole milk ½ gal, lot 24-117", "Customer complaint: foreign material", at(390, 9, 0), at(390, 14, 40), None, 4800, 4610, None, None, "Two pallets at 3PL not in system", "CA-061"])
    _put(mr, 11, [d(150), "Both", "Chocolate milk qt, lot 26-042", "Supplier recall: cocoa powder", at(150, 8, 30), at(150, 11, 5), None, 3600, 3600])
    for i, sec in enumerate(sections):
        r = 5 + i
        planned = today + dt.timedelta(days=-300 + i * 30)
        done = planned + dt.timedelta(days=rnd.randint(0, 6)) if planned < today - dt.timedelta(days=8) else None
        if i == 9:
            done = None  # one overdue audit
        vals = [None, rnd.choice(["M. Chen", "P. Walsh", "QA Manager"]), planned, done,
                rnd.randint(0, 4) if done else None, "CA-1" + str(rnd.randint(10, 13)) if done and rnd.random() < 0.4 else None]
        _put(ia, r, vals)
    products = [("Whole milk gal", "Line 1", {"Milk": "Contains"}),
                ("2% milk ½ gal", "Line 1", {"Milk": "Contains"}),
                ("Chocolate milk qt", "Line 2", {"Milk": "Contains", "Soy": "Contains"}),
                ("Almond-vanilla creamer", "Line 2", {"Milk": "Contains", "Tree nuts": "Contains"}),
                ("Egg nog qt", "Line 2", {"Milk": "Contains", "Egg": "Contains", "Tree nuts": "Shared line"}),
                ("Cookie dough ice cream", "Line 3", {"Milk": "Contains", "Egg": "Contains", "Wheat": "Contains", "Soy": "Contains"}),
                ("Vanilla ice cream", "Line 3", {"Milk": "Contains", "Egg": "Contains", "Wheat": "Shared line", "Peanuts": "Shared line"}),
                ("Peanut butter swirl", "Line 3", {"Milk": "Contains", "Peanuts": "Contains", "Soy": "Contains"})]
    for i, (p, line, al_map) in enumerate(products):
        r = 5 + i
        _put(al, r, [p, line] + [al_map.get(a) for a in ALLERGENS])
        al.cell(row=r, column=3 + len(ALLERGENS) + 2, value=d(rnd.choice([30, 90, 200, 420]) if i != 4 else 500) if i != 7 else None)
    for i, line in enumerate(["Line 1", "Line 2", "Line 3"]):
        al.cell(row=5 + i, column=3 + len(ALLERGENS) + 5, value=line)
    glass = [("G-001", "Pasteurization", "HTST panel", "Recorder chart window", "Glass", 1, "Weekly", 3, "Intact"),
             ("G-002", "QA Lab", "Lab bench", "Glass thermometers (sealed case)", "Glass", 4, "Weekly", 9, "Intact"),
             ("G-003", "Filling / Packaging", "Filler 2", "Sight glass", "Glass", 2, "Weekly", 2, "Intact"),
             ("G-004", "Filling / Packaging", "Filler room", "Light covers (shatterproof)", "Brittle plastic", 12, "Monthly", 12, "Intact"),
             ("G-005", "Warehouse", "Cooler dock", "Forklift mirror", "Acrylic", 2, "Monthly", 20, "Damaged"),
             ("G-006", "Raw Processing", "Silo room", "Tank level sight tubes", "Glass", 3, "Weekly", 10, "Intact"),
             ("G-007", "Maintenance", "Shop", "Safety glasses station", "Brittle plastic", 1, "Quarterly", 40, "Intact")]
    for i, g in enumerate(glass):
        _put(gl, 5 + i, list(g[:7]) + [d(g[7]), g[8], rnd.choice(["J. Ortiz", "M. Chen"])])
    r = 5
    for n in range(360, 0, -7):
        for dev, dtype, area, pest, rate in (("RT-04", "Rodent trap", "Interior - warehouse", "Mouse", 0.12),
                                             ("ILT-02", "Insect light trap", "Interior - processing", "Flies", 1.0),
                                             ("PH-01", "Pheromone trap", "Interior - warehouse", "Stored-product insects", 0.35),
                                             ("BS-11", "Bait station", "Exterior", "Mouse", 0.08)):
            if rnd.random() < rate:
                month = (today - dt.timedelta(days=n)).month
                count = rnd.randint(1, 4) if dtype != "Insect light trap" else rnd.randint(1, 6) + (5 if month in (7, 8) else 0)
                _put(pt, r, [d(n), dev, dtype, area, pest, count, "Removed, device serviced", "PCO"])
                r += 1
    for dev, dtype, area, pest, cnt in (("RT-04", "Rodent trap", "Interior - warehouse", "Mouse", 3),
                                        ("RT-06", "Rodent trap", "Interior - warehouse", "Mouse", 1)):
        _put(pt, r, [d(max(0, min(4, today.day - 1))), dev, dtype, area, pest, cnt, "Exclusion check on dock door seals", "PCO"])
        r += 1
    cals = [("TH-01", "HTST recording thermometer", "Pasteurization", "Recording thermometer", "Indicating thermometer", "±0.5°F", "Quarterly", 60, "Pass"),
            ("TH-07", "Digital probe thermometer", "QA Lab", "Thermometer", "NIST reference thermometer", "±1°F", "Monthly", 38, "Pass"),
            ("SC-02", "Filler check-weigh scale", "Filling / Packaging", "Scale", "Certified test weights", "±2 g", "Monthly", 12, "Adjusted"),
            ("PH-01", "Lab pH meter", "QA Lab", "pH meter", "pH 4/7/10 buffers", "±0.05", "Daily", 0, "Pass"),
            ("MD-01", "Metal detector line 3", "Filling / Packaging", "Metal detector", "Fe/NFe/SS test pieces", "Detect all", "Annual", 300, "Pass"),
            ("PG-03", "Homogenizer pressure gauge", "Pasteurization", "Pressure gauge", "Master gauge", "±2%", "Semiannual", 200, "Fail")]
    for i, c in enumerate(cals):
        _put(ca, 5 + i, list(c[:7]) + [d(c[7]), c[8], rnd.choice(["S. Brooks", "M. Chen"])])
