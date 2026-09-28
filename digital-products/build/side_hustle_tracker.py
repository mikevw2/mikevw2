"""Builds the 'Side Hustle Profit & Tax Tracker' workbook."""
import datetime as dt
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from common import (BOX, DARK, DATE, GREY, INPUT, LIGHT, MID, MONEY, PCT, RED, WHITE,
                    fill, header, section, style_range, title, widths)

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
# Categories modeled on IRS Schedule C lines so totals transfer straight to a tax return.
EXPENSE_CATS = ["Advertising", "Car & truck (actual)", "Commissions & fees", "Contract labor",
                "Insurance", "Legal & professional", "Office expense", "Rent / lease",
                "Repairs & maintenance", "Supplies", "Taxes & licenses", "Travel", "Meals (50%)",
                "Utilities", "Software & subscriptions", "Phone & internet", "Education",
                "Equipment", "Shipping & postage", "Cost of goods sold", "Other"]
PLATFORMS = ["Direct client", "Etsy", "Shopify", "Upwork", "Fiverr", "Amazon", "Uber / Lyft",
             "DoorDash", "YouTube / Ads", "Other"]
ROWS = 1000
QUARTERS = [("Q1", 1, 3, "Apr 15"), ("Q2", 4, 5, "Jun 15"), ("Q3", 6, 8, "Sep 15"), ("Q4", 9, 12, "Jan 15 (next yr)")]


def build(path, sample=False, year=None, hook=None):
    year = year or dt.date.today().year
    wb = Workbook()
    start = wb.active
    start.title = "Start Here"
    setup = wb.create_sheet("Setup")
    inc = wb.create_sheet("Income")
    exp = wb.create_sheet("Expenses")
    mil = wb.create_sheet("Mileage")
    dash = wb.create_sheet("Dashboard")
    tax = wb.create_sheet("Quarterly Taxes")

    # ---------------- Start Here ----------------
    title(start, "Side Hustle Profit & Tax Tracker", "Know your real profit — and never get surprised by a tax bill.", 10)
    steps = [
        ("1. Setup", "Enter your business name, year, estimated tax rates and the IRS mileage rate."),
        ("2. Income", "Log every sale or payout. Record platform fees separately so you see true net income."),
        ("3. Expenses", "Log business costs using categories that match IRS Schedule C. Set the business-use % for mixed items."),
        ("4. Mileage", "Log business trips; the deduction is calculated for you."),
        ("5. Dashboard", "Monthly income, expenses and profit, plus a year-to-date summary by category."),
        ("6. Quarterly Taxes", "See an estimate of what to set aside for each quarterly payment."),
    ]
    r = 4
    for h, b in steps:
        start.cell(row=r, column=2, value=h).font = Font(bold=True, size=12, color=DARK)
        c = start.cell(row=r, column=3, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        start.row_dimensions[r].height = 32
        r += 1
    r += 1
    note = start.cell(row=r, column=3, value=(
        "Important: tax figures are estimates for planning only and are not tax advice. "
        "Rates and rules change every year — confirm with the IRS or a tax professional."))
    note.font = Font(italic=True, color="9A3412")
    note.alignment = Alignment(wrap_text=True)
    start.row_dimensions[r].height = 32
    widths(start, {"B": 22, "C": 100})

    # ---------------- Setup ----------------
    title(setup, "Setup", "Yellow cells are inputs.", 6)
    fields = [("Business name", "My Side Hustle", None),
              ("Tax year", year, "0"),
              ("Estimated federal income tax rate", 0.12, PCT),
              ("Estimated state income tax rate", 0.05, PCT),
              ("Self-employment tax rate", 0.153, PCT),
              ("Share of profit subject to SE tax", 0.9235, PCT),
              ("Mileage rate ($/mile) — update yearly from irs.gov", 0.70, '"$"0.000')]
    for i, (lab, val, fmt) in enumerate(fields):
        setup.cell(row=4 + i, column=2, value=lab).font = Font(bold=True)
        c = setup.cell(row=4 + i, column=3, value=val)
        c.fill = fill(INPUT)
        c.border = BOX
        if fmt:
            c.number_format = fmt
    setup["B12"] = "Effective set-aside rate on profit"
    setup["B12"].font = Font(bold=True, color=DARK)
    setup["C12"] = "=C8*C9+(C6+C7)*(1-C8*C9/2)"
    setup["C12"].number_format = PCT
    setup["C12"].font = Font(bold=True, size=12, color=DARK)
    setup["C12"].fill = fill(LIGHT)
    setup["B13"] = "(SE tax on 92.35% of profit, plus income tax on profit after the SE-tax deduction.)"
    setup["B13"].font = Font(italic=True, color="777777")
    section(setup, "E3", "Expense categories")
    for i, c in enumerate(EXPENSE_CATS):
        cell = setup.cell(row=4 + i, column=5, value=c)
        cell.fill = fill(INPUT)
        cell.border = BOX
    section(setup, "G3", "Platforms")
    for i, p in enumerate(PLATFORMS):
        cell = setup.cell(row=4 + i, column=7, value=p)
        cell.fill = fill(INPUT)
        cell.border = BOX
    widths(setup, {"B": 48, "C": 18, "D": 3, "E": 26, "F": 3, "G": 20})
    CAT_LIST = f"=Setup!$E$4:$E${3 + len(EXPENSE_CATS) + 4}"
    PLAT_LIST = f"=Setup!$G$4:$G${3 + len(PLATFORMS) + 4}"
    Y = "Setup!$C$5"

    # ---------------- Income ----------------
    title(inc, "Income", "Gross = what the customer paid. Fees = platform/processing fees.", 9)
    header(inc, 4, 1, ["Date", "Client / order", "Platform", "Description", "Gross", "Fees", "Net", "Month", "Year"])
    inc.freeze_panes = "A5"
    dv = DataValidation(type="list", formula1=PLAT_LIST, allow_blank=True)
    inc.add_data_validation(dv)
    dv.add(f"C5:C{4 + ROWS}")
    for r in range(5, 5 + ROWS):
        inc.cell(row=r, column=1).number_format = DATE
        inc.cell(row=r, column=5).number_format = MONEY
        inc.cell(row=r, column=6).number_format = MONEY
        g = inc.cell(row=r, column=7, value=f'=IF(A{r}="","",N(E{r})-N(F{r}))')
        g.number_format = MONEY
        inc.cell(row=r, column=8, value=f'=IF(A{r}="","",MONTH(A{r}))').font = Font(color="999999")
        inc.cell(row=r, column=9, value=f'=IF(A{r}="","",YEAR(A{r}))').font = Font(color="999999")
        if r % 2 == 0:
            for c in range(1, 7):
                inc.cell(row=r, column=c).fill = fill(GREY)
    widths(inc, {"A": 14, "B": 22, "C": 16, "D": 30, "E": 12, "F": 12, "G": 12, "H": 7, "I": 7})

    # ---------------- Expenses ----------------
    title(exp, "Expenses", "Business-use % lets you deduct only the business share of mixed-use items.", 10)
    header(exp, 4, 1, ["Date", "Vendor", "Category", "Description", "Amount", "Business %", "Deductible", "Receipt?", "Month", "Year"])
    exp.freeze_panes = "A5"
    dvc = DataValidation(type="list", formula1=CAT_LIST, allow_blank=True)
    dvr = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    exp.add_data_validation(dvc)
    exp.add_data_validation(dvr)
    dvc.add(f"C5:C{4 + ROWS}")
    dvr.add(f"H5:H{4 + ROWS}")
    for r in range(5, 5 + ROWS):
        exp.cell(row=r, column=1).number_format = DATE
        exp.cell(row=r, column=5).number_format = MONEY
        exp.cell(row=r, column=6).number_format = "0%"
        d = exp.cell(row=r, column=7,
                     value=f'=IF(A{r}="","",N(E{r})*IF(F{r}="",1,F{r})*IF(C{r}="Meals (50%)",0.5,1))')
        d.number_format = MONEY
        exp.cell(row=r, column=9, value=f'=IF(A{r}="","",MONTH(A{r}))').font = Font(color="999999")
        exp.cell(row=r, column=10, value=f'=IF(A{r}="","",YEAR(A{r}))').font = Font(color="999999")
        if r % 2 == 0:
            for c in range(1, 9):
                exp.cell(row=r, column=c).fill = fill(GREY)
    exp.conditional_formatting.add(f"H5:H{4 + ROWS}", CellIsRule(operator="equal", formula=['"No"'], fill=fill(RED)))
    widths(exp, {"A": 14, "B": 20, "C": 24, "D": 28, "E": 12, "F": 11, "G": 12, "H": 10, "I": 7, "J": 7})

    # ---------------- Mileage ----------------
    title(mil, "Mileage Log", "The IRS expects date, purpose and miles for every business trip.", 8)
    header(mil, 4, 1, ["Date", "Purpose / destination", "Start odometer", "End odometer", "Miles", "Deduction", "Month", "Year"])
    mil.freeze_panes = "A5"
    for r in range(5, 5 + ROWS):
        mil.cell(row=r, column=1).number_format = DATE
        mil.cell(row=r, column=5, value=f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})')
        mil.cell(row=r, column=6, value=f'=IF(E{r}="","",E{r}*Setup!$C$10)').number_format = MONEY
        mil.cell(row=r, column=7, value=f'=IF(A{r}="","",MONTH(A{r}))').font = Font(color="999999")
        mil.cell(row=r, column=8, value=f'=IF(A{r}="","",YEAR(A{r}))').font = Font(color="999999")
    widths(mil, {"A": 14, "B": 34, "C": 15, "D": 15, "E": 10, "F": 12, "G": 7, "H": 7})

    last = 4 + ROWS
    IG, IF_, IM, IY = (f"Income!$E$5:$E${last}", f"Income!$F$5:$F${last}",
                       f"Income!$H$5:$H${last}", f"Income!$I$5:$I${last}")
    ED, EC, EM, EY = (f"Expenses!$G$5:$G${last}", f"Expenses!$C$5:$C${last}",
                      f"Expenses!$I$5:$I${last}", f"Expenses!$J$5:$J${last}")
    MD, MM, MY = f"Mileage!$F$5:$F${last}", f"Mileage!$G$5:$G${last}", f"Mileage!$H$5:$H${last}"

    # ---------------- Dashboard ----------------
    title(dash, "Dashboard", None, 11)
    dash["B2"] = '=Setup!$C$4&" — "&Setup!$C$5'
    dash["B2"].font = Font(italic=True, color="CFE3D8")
    header(dash, 4, 2, ["Month", "Gross income", "Platform fees", "Expenses", "Mileage", "Net profit", "Set aside for tax"])
    for i, m in enumerate(MONTHS):
        r = 5 + i
        n = i + 1
        dash.cell(row=r, column=2, value=m)
        dash.cell(row=r, column=3, value=f"=SUMIFS({IG},{IM},{n},{IY},{Y})")
        dash.cell(row=r, column=4, value=f"=SUMIFS({IF_},{IM},{n},{IY},{Y})")
        dash.cell(row=r, column=5, value=f"=SUMIFS({ED},{EM},{n},{EY},{Y})")
        dash.cell(row=r, column=6, value=f"=SUMIFS({MD},{MM},{n},{MY},{Y})")
        dash.cell(row=r, column=7, value=f"=C{r}-D{r}-E{r}-F{r}")
        dash.cell(row=r, column=8, value=f"=MAX(0,G{r})*Setup!$C$12")
    style_range(dash, "B5:B16")
    style_range(dash, "C5:H16", MONEY)
    dash.conditional_formatting.add("G5:G16", CellIsRule(operator="lessThan", formula=["0"], fill=fill(RED)))
    dash.cell(row=17, column=2, value="Year total")
    for c, L in enumerate("CDEFGH", start=3):
        dash.cell(row=17, column=c, value=f"=SUM({L}5:{L}16)")
    style_range(dash, "B17:H17", color=LIGHT, bold=True)
    style_range(dash, "C17:H17", MONEY, color=LIGHT, bold=True)
    dash["B19"] = "Profit margin"
    dash["C19"] = '=IF(C17=0,"",G17/C17)'
    dash["C19"].number_format = PCT
    dash["B19"].font = Font(bold=True)
    dash["C19"].font = Font(bold=True, size=12, color=DARK)

    section(dash, "J4", "Expenses by category (year)")
    header(dash, 5, 10, ["Category", "Deductible"])
    for i in range(len(EXPENSE_CATS) + 4):
        r = 6 + i
        dash.cell(row=r, column=10, value=f'=IF(Setup!E{4 + i}="","",Setup!E{4 + i})')
        dash.cell(row=r, column=11, value=f'=IF(J{r}="","",SUMIFS({ED},{EC},J{r},{EY},{Y}))')
    style_range(dash, f"J6:J{5 + len(EXPENSE_CATS) + 4}")
    style_range(dash, f"K6:K{5 + len(EXPENSE_CATS) + 4}", MONEY)

    ch = BarChart()
    ch.title = "Monthly net profit"
    ch.height = 8
    ch.width = 18
    ch.legend = None
    ch.add_data(Reference(dash, min_col=7, min_row=4, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=2, min_row=5, max_row=16))
    dash.add_chart(ch, "B21")
    widths(dash, {"B": 12, "C": 14, "D": 14, "E": 14, "F": 12, "G": 14, "H": 16, "I": 3, "J": 26, "K": 14})

    # ---------------- Quarterly Taxes ----------------
    title(tax, "Quarterly Estimated Taxes", "Estimates for planning — not tax advice.", 8)
    header(tax, 4, 2, ["Quarter", "Covers", "Due date", "Net profit", "Est. tax to set aside", "Amount paid", "Difference"])
    for i, (q, m1, m2, due) in enumerate(QUARTERS):
        r = 5 + i
        tax.cell(row=r, column=2, value=q)
        tax.cell(row=r, column=3, value=f"{MONTHS[m1 - 1]}–{MONTHS[m2 - 1]}")
        tax.cell(row=r, column=4, value=due)
        tax.cell(row=r, column=5, value=f"=SUM(Dashboard!G{4 + m1}:G{4 + m2})")
        tax.cell(row=r, column=6, value=f"=MAX(0,E{r})*Setup!$C$12")
        p = tax.cell(row=r, column=7)
        p.fill = fill(INPUT)
        tax.cell(row=r, column=8, value=f"=N(G{r})-F{r}")
    style_range(tax, "B5:D8")
    style_range(tax, "E5:H8", MONEY)
    for r in range(5, 9):
        tax.cell(row=r, column=7).fill = fill(INPUT)
    tax["B9"] = "Total"
    for L in "EFGH":
        tax[f"{L}9"] = f"=SUM({L}5:{L}8)"
    style_range(tax, "B9:H9", color=LIGHT, bold=True)
    style_range(tax, "E9:H9", MONEY, color=LIGHT, bold=True)
    tax["B11"] = ("IRS quarterly periods are uneven (Q2 is only Apr–May). Pay via IRS Direct Pay or EFTPS. "
                  "If you also have a W-2 job, raising your W-4 withholding can replace quarterly payments.")
    tax["B11"].alignment = Alignment(wrap_text=True)
    tax.merge_cells("B11:H12")
    tax.row_dimensions[11].height = 30
    widths(tax, {"B": 10, "C": 12, "D": 18, "E": 14, "F": 20, "G": 14, "H": 14})

    for ws in (setup, inc, exp, mil):
        ws.sheet_properties.tabColor = "F2C94C"
    for ws in (dash, tax):
        ws.sheet_properties.tabColor = MID
    if sample:
        _sample(inc, exp, mil, year)
    wb.active = 0
    if hook:
        hook(wb)
    wb.save(path)


def _sample(inc, exp, mil, year):
    rnd = random.Random(7)
    r = 5
    for m in range(1, 13):
        for _ in range(rnd.randint(6, 12)):
            g = round(rnd.uniform(25, 180), 2)
            vals = [dt.date(year, m, rnd.randint(1, 28)), f"Order #{rnd.randint(1000, 9999)}", "Etsy",
                    "Digital + print orders", g, round(g * 0.095 + 0.45, 2)]
            for j, v in enumerate(vals):
                inc.cell(row=r, column=1 + j, value=v)
            r += 1
        if m % 3 == 0:
            vals = [dt.date(year, m, 20), "Acme Co.", "Direct client", "Logo design project", 850, 0]
            for j, v in enumerate(vals):
                inc.cell(row=r, column=1 + j, value=v)
            r += 1
    r = 5
    for m in range(1, 13):
        items = [("Adobe", "Software & subscriptions", "Creative Cloud", 59.99, None),
                 ("Verizon", "Phone & internet", "Phone bill", 85, 0.5),
                 ("Meta", "Advertising", "Instagram ads", round(rnd.uniform(40, 150), 2), None),
                 ("USPS", "Shipping & postage", "Shipping labels", round(rnd.uniform(30, 90), 2), None),
                 ("Michaels", "Supplies", "Packaging + paper", round(rnd.uniform(20, 75), 2), None)]
        if m == 4:
            items.append(("Apple", "Equipment", "iPad for design work", 799, 1))
        if m in (5, 10):
            items.append(("Cafe", "Meals (50%)", "Client meeting", 64.2, None))
        for vendor, cat, desc, amt, pct in items:
            vals = [dt.date(year, m, rnd.randint(1, 28)), vendor, cat, desc, amt, pct, None, rnd.choice(["Yes", "Yes", "Yes", "No"])]
            for j, v in enumerate(vals):
                if j != 6:
                    exp.cell(row=r, column=1 + j, value=v)
            r += 1
    r = 5
    odo = 42150
    for m in range(1, 13):
        for _ in range(2):
            miles = rnd.randint(8, 46)
            vals = [dt.date(year, m, rnd.randint(1, 28)), "Post office / supply run", odo, odo + miles]
            for j, v in enumerate(vals):
                mil.cell(row=r, column=1 + j, value=v)
            odo += miles + rnd.randint(50, 300)
            r += 1


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    build(f"{out}/Side-Hustle-Profit-and-Tax-Tracker.xlsx", sample=False)
    build(f"{out}/Side-Hustle-Profit-and-Tax-Tracker-SAMPLE.xlsx", sample=True, year=2026)
