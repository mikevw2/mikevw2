import sys
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1],data_only=True)
d=wb["Debt Payoff"]
for r in range(5,14): print([d.cell(r,c).value for c in range(2,8)])
for r in range(10,15): print(d.cell(r,9).value, d.cell(r,10).value)
m=wb["Monthly Dashboard"]; print([m.cell(7,c).value for c in range(3,8)])
for r in range(11,16): print([m.cell(r,c).value for c in range(2,7)])
a=wb["Annual Overview"]
for row in a.iter_rows(min_row=18,max_row=20,min_col=2,max_col=16,values_only=True): print(row)
g=wb["Savings Goals"]
for row in g.iter_rows(min_row=5,max_row=8,min_col=2,max_col=9,values_only=True): print(row)
errs=[]
for ws in wb:
  for row in ws.iter_rows():
    for c in row:
      v=c.value
      if isinstance(v,str) and (v.startswith("#") or v.startswith("Err:")): errs.append((ws.title,c.coordinate,v))
print("errors",len(errs),errs[:5])
